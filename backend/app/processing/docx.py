"""DOCX processor: headings, paragraphs, tables (python-docx, local)."""
from pathlib import Path

from docx import Document as DocxDocument

from ..core.config import get_settings
from .base import DocumentProcessor, ExtractResult, ProcessingFailed, flatten_sections


class DOCXProcessor(DocumentProcessor):
    file_types = ("DOCX",)

    def process(self, path: Path, *, original_filename: str) -> ExtractResult:
        try:
            doc = DocxDocument(str(path))
        except Exception as e:
            raise ProcessingFailed("Unable to read this Word document") from e
        sections: list[dict] = []
        for para in doc.paragraphs:
            text = (para.text or "").strip()
            if not text:
                continue
            style = (para.style.name or "").lower()
            if style.startswith("heading"):
                sections.append({"type": "heading", "text": text, "style": para.style.name})
            else:
                sections.append({"type": "paragraph", "text": text})
        for idx, table in enumerate(doc.tables, start=1):
            rows = []
            for row in table.rows:
                rows.append(" | ".join((c.text or "").strip() for c in row.cells))
            body = "\n".join(r for r in rows if r.strip())
            if body:
                sections.append({"type": "table", "source": f"Table {idx}", "text": body})
        s = get_settings()
        text, truncated = flatten_sections(sections, s.EXTRACT_MAX_CHARS)
        note = ("extracted text truncated"
                f" at {s.EXTRACT_MAX_CHARS} characters") if truncated else None
        if not text:
            raise ProcessingFailed("No readable text found in this Word document")
        return ExtractResult(ok=True, extracted_text=text, sections=sections,
                             page_count=None, ocr_used=False, note=note)
