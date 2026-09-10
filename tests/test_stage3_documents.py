"""Stage 3 tests: upload/validation, hashing, storage safety, isolation, RBAC,
processing per format, OCR honesty, versioning, equipment link, download,
archive/delete, audit, security probes."""
import hashlib
import io

from helpers import client, db

K = 0


def _next(prefix="kb"):
    global K
    K += 1
    return f"{prefix}{K}@kb.test", f"{prefix.capitalize()} Co {K}"


def _admin():
    email, company = _next()
    r = client.post("/api/auth/register", json={
        "company_name": company, "name": "Admin", "email": email,
        "password": "Str0ngPass!"})
    assert r.status_code == 201, r.text
    lr = client.post("/api/auth/login",
                     json={"email": email, "password": "Str0ngPass!"})
    assert lr.status_code == 200
    return {"Authorization": f"Bearer {lr.json()['access_token']}"}, email


def _mkuser(admin_h, role, tag):
    email = f"{tag}-{role.lower()}@kb.test"
    r = client.post("/api/auth/users", headers=admin_h,
                    json={"name": tag, "email": email,
                          "password": "Str0ngPass!", "role": role})
    assert r.status_code == 201, r.text
    lr = client.post("/api/auth/login",
                     json={"email": email, "password": "Str0ngPass!"})
    assert lr.status_code == 200
    return {"Authorization": f"Bearer {lr.json()['access_token']}"}


def _equipment(admin_h, code="P-204"):
    r = client.post("/api/equipment", headers=admin_h,
                    json={"code": code, "name": f"Pump {code}", "type": "Pump",
                          "criticality": "HIGH"})
    assert r.status_code == 201, r.text
    return r.json()["id"]


def _upload(h, content: bytes, filename: str, equipment_id=None,
            mime="application/octet-stream"):
    files = {"file": (filename, io.BytesIO(content), mime)}
    data = {}
    if equipment_id is not None:
        data["equipment_id"] = str(equipment_id)
    return client.post("/api/documents", headers=h, files=files, data=data)


def _pdf_bytes(text="Vibration report P-204"):
    # Minimal valid single-page PDF with real text (offsets computed).
    objs = []
    objs.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objs.append(b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>")
    objs.append(b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 200]"
                b" /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>")
    stream = f"BT /F1 14 Tf 50 150 Td ({text}) Tj ET".encode("latin-1")
    objs.append(b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n"
                + stream + b"\nendstream")
    objs.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objs, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs)+1}\n0000000000 65535 f \n".encode()
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += (f"trailer\n<< /Size {len(objs)+1} /Root 1 0 R >>\n"
            f"startxref\n{xref}\n%%EOF").encode()
    return bytes(out)


def _docx_bytes(heading="Pump Manual", para="Torque bolts."):
    from docx import Document
    import io as _io
    d = Document()
    d.add_heading(heading, 1)
    d.add_paragraph(para)
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "Part"
    t.cell(0, 1).text = "Nm"
    t.cell(1, 0).text = "Bolt"
    t.cell(1, 1).text = "40"
    buf = _io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _xlsx_bytes():
    import openpyxl
    import io as _io
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Readings"
    ws.append(["Date", "Temp"])
    ws.append(["2026-09-01", 72.4])
    buf = _io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _png_bytes(color="red", size=(120, 80)):
    from PIL import Image
    import io as _io
    buf = _io.BytesIO()
    Image.new("RGB", size, color).save(buf, format="PNG")
    return buf.getvalue()


def test_auth_required():
    assert client.post("/api/documents").status_code in (401, 422)
    assert client.get("/api/documents").status_code == 401
    assert client.get("/api/documents/1").status_code == 401
    assert client.get("/api/documents/1/download").status_code == 401


def test_upload_all_formats_and_extraction():
    ha, _ = _admin()
    cases = [
        (_pdf_bytes(), "report.pdf", "PDF", "Vibration report P-204", 1),
        (_docx_bytes(), "manual.docx", "DOCX", "Pump Manual", None),
        (b"Check bearings monthly.\nSecond line.", "notes.txt", "TXT",
         "bearings", None),
        (b"timestamp,temperature\n2026-09-01 10:00,72.4\n", "s.csv", "CSV",
         "temperature=72.4", None),
        (_xlsx_bytes(), "r.xlsx", "XLSX", "Readings", None),
        (_png_bytes(), "pump.png", "PNG", None, None),
    ]
    for content, name, ftype, needle, pages in cases:
        r = _upload(ha, content, name)
        assert r.status_code == 201, (name, r.text)
        d = r.json()
        assert d["file_type"] == ftype, name
        assert d["sha256_hash"] == hashlib.sha256(content).hexdigest(), name
        assert d["processing_status"] == "COMPLETED", (name, d)
        assert d["ocr_used"] is False, name  # no tesseract in this env
        if needle:
            assert needle in d["text_preview"] or _full_text(ha, d["id"], needle), name
        if pages:
            assert d["page_count"] == pages, name
    # DOCX structure preserved for Stage 4
    r = _upload(ha, _docx_bytes(), "struct.docx")
    prev = client.get(f"/api/documents/{r.json()['id']}/preview", headers=ha).json()
    types = [s["type"] for s in prev["sections"]]
    assert "heading" in types and "table" in types


def _full_text(h, doc_id, needle):
    p = client.get(f"/api/documents/{doc_id}/preview", headers=h).json()
    return needle in p["extracted_text"]


def test_validation_rejections():
    ha, _ = _admin()
    assert _upload(ha, b"MZ...", "evil.exe").status_code == 422
    assert _upload(ha, b"hello", "nofile").status_code == 422
    assert _upload(ha, b"", "empty.txt").status_code == 422
    assert _upload(ha, _png_bytes(), "../escape.pdf").status_code == 422
    assert _upload(ha, _png_bytes(), "fake.pdf").status_code == 422  # MIME mismatch
    assert _upload(ha, b"<script>alert(1)</script>", "x.txt").status_code == 201
    # Real-world vendor MIME variants are accepted (content still sniffed).
    r = client.post("/api/documents", headers=ha,
                    files={"file": ("r.csv", io.BytesIO(b"t,v\n1,2\n"),
                                    "application/vnd.ms-excel")})
    assert r.status_code == 201, r.text
    # Definite contradictions still rejected.
    r = client.post("/api/documents", headers=ha,
                    files={"file": ("r2.csv", io.BytesIO(b"t,v\n1,2\n"),
                                    "image/png")})
    assert r.status_code == 422


def test_oversize_rejected():
    from app.core.config import get_settings
    s = get_settings()
    old = s.MAX_UPLOAD_SIZE_MB
    s.MAX_UPLOAD_SIZE_MB = 0
    try:
        ha, _ = _admin()
        r = _upload(ha, b"tiny but over limit", "t.txt")
        assert r.status_code == 413
        assert "limit" in r.json()["detail"].lower()
    finally:
        s.MAX_UPLOAD_SIZE_MB = old


def test_malformed_pdf_fails_safely():
    ha, _ = _admin()
    bad = b"%PDF-1.4\n" + b"\x00\xff garbage" * 100
    r = _upload(ha, bad, "broken.pdf")
    assert r.status_code == 201
    d = r.json()
    assert d["processing_status"] == "FAILED"
    assert d["processing_error"]
    for leak in ("Traceback", "site-packages", "C:\\", ".py"):
        assert leak not in d["processing_error"]
    # retry keeps same version, stays honest
    r2 = client.post(f"/api/documents/{d['id']}/retry", headers=ha)
    assert r2.status_code == 200
    assert r2.json()["version"] == d["version"]
    assert r2.json()["processing_status"] == "FAILED"
    # retry of a COMPLETED doc is refused
    ok = _upload(ha, b"fine", "ok.txt").json()
    assert client.post(
        f"/api/documents/{ok['id']}/retry", headers=ha).status_code == 409


def test_retry_success_path():
    ha, _ = _admin()
    d = _upload(ha, b"retry me", "r.txt").json()
    con = db()
    try:
        con.execute("UPDATE documents SET processing_status = 'FAILED',"
                    " processing_error = 'simulated' WHERE id = ?", (d["id"],))
        con.commit()
    finally:
        con.close()
    r = client.post(f"/api/documents/{d['id']}/retry", headers=ha)
    assert r.status_code == 200
    assert r.json()["processing_status"] == "COMPLETED"
    assert r.json()["version"] == d["version"]  # retry never versions


def test_hash_duplicate_and_versioning():
    ha, _ = _admin()
    eq = _equipment(ha)
    a = _upload(ha, _pdf_bytes("manual v1"), "Manual.pdf", equipment_id=eq).json()
    assert a["version"] == 1 and a["parent_document_id"] is None
    dup = _upload(ha, _pdf_bytes("manual v1"), "Manual.pdf", equipment_id=eq)
    assert dup.status_code == 200 and dup.json().get("duplicate") is True
    assert dup.json()["id"] == a["id"]  # no accidental version
    b = _upload(ha, _pdf_bytes("manual v2 content"), "manual.PDF",
                equipment_id=eq).json()
    assert b["version"] == 2 and b["parent_document_id"] == a["id"]
    # v1 retained and retrievable
    assert client.get(f"/api/documents/{a['id']}", headers=ha).status_code == 200
    # same name, different equipment → separate chain
    eq2 = _equipment(ha, code="P-205")
    c = _upload(ha, _pdf_bytes("manual v1"), "Manual.pdf", equipment_id=eq2).json()
    assert c["version"] == 1


def test_storage_layout_and_safety():
    ha, _ = _admin()
    content = b"storage check"
    d = _upload(ha, content, "s.txt").json()
    con = db()
    try:
        row = con.execute("SELECT storage_key FROM documents WHERE id = ?",
                          (d["id"],)).fetchone()
    finally:
        con.close()
    key = row["storage_key"]
    assert ".." not in key and not key.startswith("/") and ":" not in key
    from pathlib import Path
    from app.core.config import ROOT, get_settings
    p = (Path(get_settings().STORAGE_DOCUMENTS)
         if not Path(get_settings().STORAGE_DOCUMENTS).is_absolute()
         else Path(get_settings().STORAGE_DOCUMENTS))
    base = p if p.is_absolute() else ROOT / p
    target = (base / key).resolve()
    assert base.resolve() in target.parents and target.is_file()
    assert target.read_bytes() == content
    assert d["sha256_hash"] == hashlib.sha256(content).hexdigest()


def test_company_isolation():
    ha, _ = _admin()
    hb, _ = _admin()
    eq_a = _equipment(ha)
    d = _upload(ha, _pdf_bytes("secret pump data"), "secret.pdf",
                equipment_id=eq_a).json()
    did = d["id"]
    for method, path in [("get", f"/api/documents/{did}"),
                         ("get", f"/api/documents/{did}/preview"),
                         ("get", f"/api/documents/{did}/download"),
                         ("post", f"/api/documents/{did}/retry"),
                         ("post", f"/api/documents/{did}/archive"),
                         ("delete", f"/api/documents/{did}")]:
        r = client.request(method, path, headers=hb)
        assert r.status_code == 404, (method, path, r.status_code, r.text)
        assert "secret" not in r.text.lower()
    assert client.get("/api/documents", headers=hb).json()["documents"] == []
    # cross-company equipment association rejected
    eq_b = _equipment(hb)
    r = _upload(hb, b"x", "x.txt", equipment_id=eq_a)
    assert r.status_code == 404
    ok = _upload(hb, b"x", "x.txt", equipment_id=eq_b)
    assert ok.status_code == 201


def test_rbac_matrix_documents():
    ha, _ = _admin()
    he = _mkuser(ha, "ENGINEER", "rbace")
    ht = _mkuser(ha, "TECHNICIAN", "rbact")
    # technician may upload + view + download ...
    t = _upload(ht, b"field note", "field.txt")
    assert t.status_code == 201, t.text
    did = t.json()["id"]
    assert client.get(f"/api/documents/{did}", headers=ht).status_code == 200
    assert client.get(f"/api/documents/{did}/download", headers=ht).status_code == 200
    # ... but may not retry / archive / delete
    assert client.post(f"/api/documents/{did}/retry", headers=ht).status_code == 403
    assert client.post(f"/api/documents/{did}/archive", headers=ht).status_code == 403
    assert client.delete(f"/api/documents/{did}", headers=ht).status_code == 403
    # engineer may retry but not archive/delete
    bad = _upload(he, b"%PDF-1.4\n\x00broken", "bad.pdf").json()
    assert bad["processing_status"] == "FAILED"
    assert client.post(
        f"/api/documents/{bad['id']}/retry", headers=he).status_code == 200
    assert client.post(
        f"/api/documents/{did}/archive", headers=he).status_code == 403
    assert client.delete(f"/api/documents/{did}", headers=he).status_code == 403
    # admin archives + deletes
    assert client.post(f"/api/documents/{did}/archive", headers=ha).status_code == 200
    remaining = client.get("/api/documents", headers=ha).json()["documents"]
    assert did not in [x["id"] for x in remaining]
    arch = client.get("/api/documents?is_archived=true", headers=ha).json()
    assert any(x["id"] == did for x in arch["documents"])
    assert client.delete(f"/api/documents/{did}", headers=ha).status_code == 200
    assert client.get(f"/api/documents/{did}", headers=ha).status_code == 404


def test_download_and_preview():
    ha, _ = _admin()
    content = _pdf_bytes("download me")
    d = _upload(ha, content, "dl.pdf").json()
    r = client.get(f"/api/documents/{d['id']}/download", headers=ha)
    assert r.status_code == 200
    assert r.headers["content-type"] == "application/pdf"
    assert r.content == content
    assert "document-" in r.headers["content-disposition"]
    p = client.get(f"/api/documents/{d['id']}/preview", headers=ha).json()
    assert "download me" in p["extracted_text"]
    assert isinstance(p["sections"], list) and p["sections"]
    # list hides full text and absolute paths
    lst = client.get("/api/documents", headers=ha).json()["documents"][0]
    assert "extracted_text" not in lst and "storage_key" not in lst
    assert "C:\\" not in (lst.get("processing_error") or "")


def test_filters_search_and_pagination():
    ha, _ = _admin()
    eq = _equipment(ha)
    _upload(ha, b"a", "alpha.txt", equipment_id=eq)
    _upload(ha, b"b", "beta.txt")
    _upload(ha, _png_bytes(), "gamma.png", equipment_id=eq)
    all_docs = client.get("/api/documents", headers=ha).json()
    assert all_docs["total"] == 3
    assert client.get("/api/documents?file_type=PNG", headers=ha).json()["total"] == 1
    assert client.get(f"/api/documents?equipment_id={eq}", headers=ha).json()["total"] == 2
    assert client.get("/api/documents?search=alpha", headers=ha).json()["total"] == 1
    assert client.get("/api/documents?processing_status=COMPLETED",
                      headers=ha).json()["total"] == 3
    p1 = client.get("/api/documents?page=1&page_size=2", headers=ha).json()
    p2 = client.get("/api/documents?page=2&page_size=2", headers=ha).json()
    assert len(p1["documents"]) == 2 and len(p2["documents"]) == 1
    # SQL-injection probe is just a harmless string
    inj = client.get("/api/documents?search=' OR '1'='1", headers=ha).json()
    assert inj["total"] == 0
    assert client.get("/api/documents?file_type=EXE", headers=ha).status_code == 422


def test_audit_trail_documents():
    ha, _ = _admin()
    eq = _equipment(ha)
    d = _upload(ha, b"audit me", "a.txt", equipment_id=eq).json()
    client.get(f"/api/documents/{d['id']}/download", headers=ha)
    client.post(f"/api/documents/{d['id']}/archive", headers=ha)
    con = db()
    try:
        rows = con.execute(
            "SELECT action, entity_type, entity_id, detail FROM audit_events"
            " WHERE entity_type = 'document' AND entity_id = ?"
            " ORDER BY id", (d["id"],)).fetchall()
    finally:
        con.close()
    actions = [r["action"] for r in rows]
    for expected in ("document_uploaded", "document_processing_started",
                     "document_processing_completed", "document_equipment_associated",
                     "document_downloaded", "document_archived"):
        assert expected in actions, actions
    assert all(r["entity_id"] == d["id"] for r in rows)
    assert all("audit me" not in (r["detail"] or "") for r in rows)


def test_ocr_abstraction_honest():
    from app.processing.ocr import get_ocr_provider, ocr_available
    prov = get_ocr_provider()
    assert prov.name == "tesseract"
    assert ocr_available() is False  # no tesseract binary in this env
    ha, _ = _admin()
    d = _upload(ha, _png_bytes(), "img.png").json()
    assert d["ocr_used"] is False
    assert "OCR" in (d["processing_note"] or "")

